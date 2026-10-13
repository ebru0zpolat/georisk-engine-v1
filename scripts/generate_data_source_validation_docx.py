#!/usr/bin/env python3
"""Generate docs/GeoRisk_DataSource_Validation_Package.docx from the live registry."""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.core.data_source_registry import DATA_SOURCE_REGISTRY
from app.core.data_source_validation import VALIDATION_CHECKS, validate_all_sources
from app.services.data_source_validation_service import build_validation_report

OUTPUT = REPO_ROOT / "docs" / "GeoRisk_DataSource_Validation_Package.docx"


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    doc.add_paragraph()


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_heading("GeoRisk Engine", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph("Data Source Validation Package")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Version v1.0 · {date.today().isoformat()} · Internal / audit-ready draft").italic = True

    doc.add_paragraph()
    doc.add_paragraph(
        "This document defines acceptance criteria and runtime validation posture for every "
        "external data plane GeoRisk Engine depends on. Satellite and climate proxies are not "
        "regulatory proof — validation confirms connector discipline and documented QA policy."
    )

    _add_heading(doc, "1. Scope and disclaimer", 1)
    doc.add_paragraph(
        "Validation checks confirm registry alignment, fallback discipline, and sector acceptance "
        "standards. They do not certify emissions against facility meters or TSRS disclosures."
    )
    doc.add_paragraph(
        "Canonical code: app/core/data_source_registry.py, app/core/data_source_validation.py"
    )
    doc.add_paragraph("API: GET /v1/data-sources/validation")

    _add_heading(doc, "2. Delivery modes", 1)
    _add_table(
        doc,
        ["Mode", "Meaning"],
        [
            ["live", "Credentials configured; CDS / Earth Engine may be called"],
            ["fallback", "Production API deterministic stub (CI / local dev)"],
            ["mock", "Audit script offline mode (AUDIT_USE_MOCK=true)"],
            ["stub", "Roadmap layer; deterministic agriculture proxies"],
        ],
    )

    report = build_validation_report()
    _add_heading(doc, "3. Validation summary", 1)
    doc.add_paragraph(
        f"Sources: {report.source_count} · Passed: {report.passed_count} · "
        f"Package version: {report.package_version}"
    )

    _add_heading(doc, "4. Per-source validation matrix", 1)
    for entry in DATA_SOURCE_REGISTRY:
        source_result = next(s for s in report.sources if s.source_id == entry.id)
        _add_heading(doc, f"{entry.name} ({entry.id})", 2)
        p = doc.add_paragraph()
        p.add_run(f"Tier: {entry.tier} · Delivery plane: {entry.delivery} · ")
        p.add_run(f"LOI: {entry.loi_status} · Runtime mode: {source_result.delivery_mode}")

        if entry.dataset_id:
            doc.add_paragraph(f"Dataset: {entry.dataset_id}")
        if entry.standards_ref:
            doc.add_paragraph(f"Standard: {entry.standards_ref}")
        doc.add_paragraph(f"Fallback policy: {entry.fallback_policy}")

        checks = VALIDATION_CHECKS.get(entry.id, ())
        if checks:
            rows = []
            for check in checks:
                result = next(c for c in source_result.checks if c.id == check.id)
                rows.append([check.id, check.description, result.status, result.detail])
            _add_table(doc, ["Check id", "Description", "Status", "Detail"], rows)
        else:
            doc.add_paragraph("No automated checks registered.")

    _add_heading(doc, "5. Integration checklist", 1)
    for item in [
        "Add or update entry in app/core/data_source_registry.py",
        "Add acceptance checks to VALIDATION_CHECKS in app/core/data_source_validation.py",
        "Import dataset constants in connector or audit script",
        "Document sector rules under docs/standards/ when applicable",
        "Expose via GET /v1/data-sources and GET /v1/data-sources/validation",
        "Regenerate this DOCX after registry changes",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    _add_heading(doc, "6. Related documents", 1)
    for link in [
        "docs/georisk_data_source_loi_summary.md — LOI / partnership posture",
        "docs/georisk_data_source_validation_summary.md — executive summary",
        "docs/standards/README.md — sector acceptance index",
        "docs/turkiye/cco-audit-methodology.md — mask policy and legal framing",
    ]:
        doc.add_paragraph(link, style="List Bullet")

    return doc


def main() -> None:
    # Touch validation to ensure imports resolve before doc build
    validate_all_sources()
    doc = build_document()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
